"""Load SeeThrough labels only after a ranking seal and resolve identities."""
import argparse
import csv
import json
import math
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import requests
from openpyxl import load_workbook
if __package__:
    from .benchmark_artifacts import benchmark_directory
else:
    from benchmark_artifacts import benchmark_directory

ROOT = next(p for p in Path(__file__).resolve().parents if (p / "核心系统").is_dir())


def numeric(value):
    return isinstance(value, (int, float)) and math.isfinite(value)


def resolve(row):
    cas = row[3]
    url = f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/name/{cas}/property/CanonicalSMILES,IsomericSMILES,Title/JSON"
    try:
        response = requests.get(url, timeout=(10, 30))
        response.raise_for_status()
        prop = response.json()["PropertyTable"]["Properties"][0]
        smiles = prop.get("ConnectivitySMILES") or prop.get("CanonicalSMILES") or prop.get("SMILES")
        return row[0], {"smiles": smiles, "pubchem_cid": prop.get("CID"), "pubchem_title": prop.get("Title"),
                        "identity_resolution": "PubChem_by_CAS", "identity_source_url": url, "identity_error": ""}
    except Exception as error:
        return row[0], {"smiles": "", "pubchem_cid": None, "pubchem_title": "",
                        "identity_resolution": "unresolved", "identity_source_url": url,
                        "identity_error": f"{type(error).__name__}: {error}"}


def main(out):
    out = benchmark_directory(out, ROOT)
    if not (out / "ranking_frozen.json").exists():
        raise RuntimeError("Reference labels are locked until ranking_frozen.json exists")
    workbook = ROOT / "数据" / "论文原始数据" / "SeeThrough补充数据2_1619个水相候选.xlsx"
    sheet = load_workbook(workbook, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(min_row=3, values_only=True))
    ba = [r for r in rows if str(r[3]).strip() == "100-51-6"]
    if len(ba) != 1:
        raise ValueError(f"Expected one BA row, found {len(ba)}")
    ba = ba[0]
    set41 = [r for r in rows if all(numeric(r[i]) for i in (4, 6, 8, 9, 10))
             and r[4] >= -1.5 and r[4] > ba[4] and r[6] > 1.58]
    set20 = [r for r in set41 if (4*(r[8]-ba[8])**2 + (r[9]-ba[9])**2 + (r[10]-ba[10])**2)**0.5 < 10]
    if (len(set41), len(set20)) != (41, 20):
        raise ValueError(f"Reference extraction mismatch: {len(set41)}, {len(set20)}")
    all_rows = {r[0]: r for r in set41}
    with (ROOT / "数据" / "软件数据库" / "补充数据3结构映射表.csv").open(encoding="utf-8-sig") as f:
        final_map = {r["论文候选编号"]: r for r in csv.DictReader(f) if r["论文候选编号"].startswith("#")}
    with (ROOT / "数据" / "软件数据库" / "补充数据2论文最终10标签表.csv").open(encoding="utf-8-sig") as f:
        final_ids = [r["论文候选编号"] for r in csv.DictReader(f) if r["论文最终10候选标签"] == "是"]
    if len(final_ids) != 10 or any(x not in all_rows for x in final_ids):
        raise ValueError("Final10 labels do not map uniquely into the frozen 41 set")
    resolved = {}
    with ThreadPoolExecutor(max_workers=8) as executor:
        jobs = [executor.submit(resolve, r) for r in set41 if r[0] not in final_map]
        for job in as_completed(jobs):
            key, value = job.result()
            resolved[key] = value
    for key, item in final_map.items():
        if key in all_rows:
            resolved[key] = {"smiles": item["SMILES"], "pubchem_cid": item["PubChemCID"],
                "pubchem_title": item["PubChem名称"], "identity_resolution": "local_verified_PubChem_mapping",
                "identity_source_url": item["结构来源URL"], "identity_error": item["错误原因"]}
    definitions = [("SeeThrough_41", set41), ("SeeThrough_20", set20),
                   ("SeeThrough_final10", [all_rows[i] for i in final_ids])]
    output = []
    for set_name, members in definitions:
        for r in members:
            output.append({"set": set_name, "reference_id": r[0], "catalog_number": r[1],
                "name": r[2], "cas": r[3], "hydration_score": r[4], "eri": r[6],
                "hsp_dD": r[8], "hsp_dP": r[9], "hsp_dH": r[10], **resolved[r[0]]})
    target = out / "seethrough_references_resolved.json"
    target.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"41": len(set41), "20": len(set20), "final10": len(final_ids),
                      "resolved_unique": sum(bool(resolved[x]["smiles"]) for x in resolved),
                      "unresolved_unique": sum(not bool(resolved[x]["smiles"]) for x in resolved)}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("out", type=Path)
    args = parser.parse_args()
    main(args.out.resolve())
