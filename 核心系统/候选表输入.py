"""候选上传输入边界：只校验并解析 CSV/XLSX，不映射业务字段或保存数据。"""

import csv
from io import BytesIO, StringIO
from pathlib import PureWindowsPath
from typing import Any

import pandas as pd
from openpyxl import load_workbook

最大上传字节数 = 2 * 1024 * 1024
支持扩展名 = {".csv", ".xlsx"}


def 安全文件名(文件名: str) -> str:
    原始 = str(文件名 or "")
    标准化 = 原始.replace("\\", "/")
    if not 标准化 or "/" in 标准化 or ":" in 标准化 or ".." in 标准化.split("/"):
        raise ValueError("上传文件名必须为不含路径、盘符或父目录的单一文件名")
    名称 = PureWindowsPath(标准化).name
    if 名称 != 标准化:
        raise ValueError("上传文件名不安全")
    return 名称

def _检查扩展名与签名(文件名: str, 内容: bytes) -> str:
    安全名称 = 安全文件名(文件名)
    后缀 = PureWindowsPath(安全名称).suffix.lower()
    if 后缀 not in 支持扩展名:
        raise ValueError("仅支持 CSV 或 XLSX 文件")
    if not 内容:
        raise ValueError("上传文件为空")
    if len(内容) > 最大上传字节数:
        raise ValueError(f"上传文件超过大小上限：{最大上传字节数} 字节")
    if 后缀 == ".xlsx" and not 内容.startswith(b"PK\x03\x04"):
        raise ValueError("XLSX 扩展名与实际文件签名不匹配")
    if 后缀 == ".csv":
        if 内容.startswith(b"PK\x03\x04") or b"\x00" in 内容[:4096]:
            raise ValueError("CSV 扩展名与实际文件内容不匹配")
        try:
            内容.decode("utf-8-sig")
        except UnicodeDecodeError as 错误:
            raise ValueError("CSV 必须使用 UTF-8 或 UTF-8-SIG 编码") from 错误
    return 后缀

def _检查重复列(列名: list[Any]) -> None:
    标准列名 = [("" if 列 is None else str(列).strip()) for 列 in 列名]
    重复 = sorted({列 for 列 in 标准列名 if 标准列名.count(列) > 1})
    if 重复:
        raise ValueError(f"上传表含重复列名，不能进入流程：{重复}")
    if any(not 列 for 列 in 标准列名):
        raise ValueError("上传表存在空列名")

def 读取候选上传内容(文件名: str, 内容: bytes) -> pd.DataFrame:
    后缀 = _检查扩展名与签名(文件名, 内容)
    try:
        if 后缀 == ".csv":
            文本 = 内容.decode("utf-8-sig")
            读取器 = csv.reader(StringIO(文本))
            首行 = next(读取器, None)
            if 首行 is None:
                raise ValueError("CSV 没有表头")
            _检查重复列(首行)
            表格 = pd.read_csv(StringIO(文本), dtype=object, keep_default_na=False)
        else:
            工作簿 = load_workbook(BytesIO(内容), read_only=True, data_only=True)
            try:
                工作表 = 工作簿.active
                首行 = [单元格.value for 单元格 in next(工作表.iter_rows(max_row=1), ())]
                _检查重复列(首行)
            finally:
                工作簿.close()
            表格 = pd.read_excel(BytesIO(内容), dtype=object, keep_default_na=False)
    except (ValueError, OSError, pd.errors.ParserError) as 错误:
        raise ValueError(f"无法读取上传文件：{错误}") from 错误
    if 表格.empty:
        raise ValueError("上传表没有数据行")
    表格.columns = [str(列).strip() for 列 in 表格.columns]
    return 表格
